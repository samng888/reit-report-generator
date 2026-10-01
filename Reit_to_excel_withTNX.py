import yfinance as yf
import pandas as pd
from datetime import datetime
import requests
from bs4 import BeautifulSoup
import os
import gc

# --- EXCEL & GRAPHICS DEPENDENCIES ---
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as OpenpyxlImage
import matplotlib.pyplot as plt

# REIT Tickers
tickers = [
    "C38U.SI", "A17U.SI", "M44U.SI", "ME8U.SI", "K71U.SI", 
    "AJBU.SI", "N2IU.SI", "J69U.SI", "T82U.SI", "C2PU.SI", 
    "0823.HK", "CMOU.SI", "AW9U.SI", "CFA.SI"
]

def get_risk_free_rate():
    """Benchmark engine: Tries US 10Y (^TNX), scrapes SG 10Y, then safe fallback."""
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }
    
    print("Fetching Global Benchmark (US 10Y: ^TNX)...")
    try:
        us_bond = yf.Ticker("^TNX")
        us_yield = us_bond.info.get('currentPrice') or us_bond.info.get('previousClose')
        if us_yield and us_yield > 0:
            print(f"-> US 10Y Rate: {round(us_yield, 2)}%")
            return float(us_yield)
    except Exception as e:
        print(f"-> US Treasury fetch bypassed: {e}")

    print("Scraping SG 10-Year Bond Yield...")
    try:
        url = "https://tradingeconomics.com/singapore/government-bond-yield"
        res = requests.get(url, headers=headers, timeout=10)
        if res.status_code == 200:
            soup = BeautifulSoup(res.text, 'html.parser')
            val = soup.find("div", {"id": "market_value"})
            if val:
                sg_yield = float(val.text.strip())
                print(f"-> Scraped SG 10Y Yield: {sg_yield}%")
                return sg_yield
    except Exception as e:
        print(f"-> Scraping fallback failed: {e}")

    print("-> Defaulting safely to baseline 3.00%.")
    return 3.00

def compute_dynamic_icr(reit):
    """Computes ICR using Quarterly Income Statement: Operating Income / Interest Expense."""
    try:
        inc_stmt = reit.quarterly_income_stmt
        if not inc_stmt.empty:
            latest = inc_stmt.iloc[:, 0]
            
            # Extract EBIT or Operating Income
            ebit = latest.get('Operating Income') or latest.get('EBIT') or latest.get('Pretax Income')
            
            # Extract Interest Expense or Finance Costs
            interest_exp = (
                latest.get('Interest Expense') or 
                latest.get('Net Non Operating Interest Income Expense') or
                latest.get('Interest Expense Non Operating')
            )
            
            if ebit is not None and interest_exp is not None and interest_exp != 0:
                icr_val = abs(float(ebit) / float(interest_exp))
                return round(icr_val, 2)
    except Exception as e:
        print(f"-> Dynamic ICR calculation bypass: {e}")
    return None

def get_reit_data():
    results = []
    current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    //output_filename = "REIT_Analysis_Dashboard.xlsx"
    
    risk_free_rate = get_risk_free_rate()
    dynamic_yield_threshold = risk_free_rate + 2.5
    
    print("\n--- Batch Downloading Historical Price Data ---")
    hist_batch = yf.download(tickers, period="1y", group_by="ticker", progress=False)

    wb = Workbook()
    ws_dashboard = wb.active
    ws_dashboard.title = "Summary Dashboard"
    temp_images = []

    for symbol in tickers:
        print(f"\nProcessing data engine for {symbol}...")
        try:
            reit = yf.Ticker(symbol)
            info = reit.info
            
            if not info or 'longName' not in info:
                continue
            
            reit_name = info.get('longName', 'Unknown Name')
            currency = info.get('currency', '???')
            price = info.get('currentPrice') or info.get('previousClose') or 0
            if price == 0:
                continue
            
            nav = info.get('bookValue', 1) or 1
            if nav <= 0: nav = 1
            p_nav = price / nav

            # --- ACCURATE TTM YIELD & PAYOUT FREQUENCY CHECK ---
            payout_status = "OK"
            try:
                dividends = reit.dividends
                if not dividends.empty:
                    one_year_ago = pd.Timestamp.now(tz=dividends.index.tz) - pd.Timedelta(days=365)
                    recent_divs = dividends[dividends.index >= one_year_ago]
                    payout_count = len(recent_divs)
                    
                    if payout_count < 2:
                        payout_status = f"FLAG: Only {payout_count} Payout(s)"
                    
                    total_dpu = recent_divs.sum()
                    div_yield = (total_dpu / price) * 100 if price > 0 else 0.0
                else:
                    div_yield = 0.0
                    payout_status = "FLAG: No Dividends"
            except Exception as e:
                print(f"-> Dividend calculation error for {symbol}: {e}")
                div_yield = 0.0
                payout_status = "ERROR"

            # --- GEARING & DYNAMIC ICR EVALUATION ---
            gearing = 0.0
            gearing_calculated = False
            d_e = info.get('debtToEquity', None)
            if d_e is not None and d_e > 0:
                gearing = (d_e / 100 / (1 + d_e / 100)) * 100
                gearing_calculated = True

            icr = info.get('interestCoverage', None)
            if icr is None or pd.isna(icr):
                icr = compute_dynamic_icr(reit)
            
            # --- DYNAMIC ICR & GEARING RISK TIERING (CONSERVATIVE FALLBACK) ---
            if gearing_calculated and gearing > 45:
                debt_tier = "HIGH RISK (Gearing > 45%)"
            elif icr is not None and icr < 1.8:
                debt_tier = "HIGH RISK (ICR < 1.8x)"
            elif icr is None:
                debt_tier = "MODERATE (Watch Zone - Missing ICR)"
            elif (40 <= gearing <= 45) or (icr is not None and icr < 3.0):
                debt_tier = "MODERATE (Watch Zone)"
            else:
                debt_tier = "HEALTHY"

            # --- TECHNICAL ANALYSIS (FROM BATCH) ---
            ema_50_val, ema_200_val, chart_img_path = None, None, None
            try:
                symbol_hist = hist_batch[symbol] if len(tickers) > 1 else hist_batch
                close_prices = symbol_hist['Close'].dropna().squeeze()
                
                if len(close_prices) >= 50:
                    ema_short = close_prices.ewm(span=50, adjust=False).mean()
                    ema_long = close_prices.ewm(span=200, adjust=False).mean()

                    ema_50_val = round(float(ema_short.iloc[-1]), 3)
                    ema_200_val = round(float(ema_long.iloc[-1]), 3)

                    # Plot Chart
                    fig, ax = plt.subplots(figsize=(10, 4.5))
                    ax.plot(close_prices.index, close_prices, label='Close Price', color='#2b2b2b', linewidth=1.5)
                    ax.plot(ema_short.index, ema_short, label=f'50 EMA: {ema_50_val}', color='#007acc', linestyle='--')
                    ax.plot(ema_long.index, ema_long, label=f'200 EMA: {ema_200_val}', color='#e056fd')
                    ax.set_title(f"{reit_name} ({symbol}) - Technical View", fontweight='bold')
                    ax.legend(loc='best')
                    ax.grid(True, alpha=0.3, linestyle=':')
                    
                    chart_img_path = f"temp_{symbol.replace('.', '_')}.png"
                    plt.tight_layout()
                    plt.savefig(chart_img_path, dpi=100)
                    plt.close(fig)
            except Exception as e:
                print(f"-> Technical processing failed for {symbol}: {e}")

            # --- HYBRID TRADING SIGNAL ENGINE ---
            is_cheap = (p_nav < 0.95) and (div_yield > dynamic_yield_threshold) and (debt_tier == "HEALTHY")
            is_expensive_or_risky = (p_nav > 1.15) or (debt_tier.startswith("HIGH RISK"))
            
            has_tech_support = (ema_50_val and price >= ema_50_val) or (ema_200_val and price >= ema_200_val)
            is_below_200_ema = (ema_200_val and price < ema_200_val)
            is_tech_breakdown = is_below_200_ema and (ema_50_val and ema_50_val < ema_200_val)

            if is_cheap and has_tech_support:
                signal = "STRONG ACCUMULATE"
            elif is_cheap:
                signal = "ACCUMULATE"
            elif is_expensive_or_risky and is_below_200_ema:
                signal = "STRONG DECUMULATE"
            elif is_expensive_or_risky or is_tech_breakdown:
                signal = "DECUMULATE"
            else:
                signal = "HOLD"

            reit_metrics = {
                "Name": reit_name, "Ticker": symbol, "Currency": currency,
                "Price": round(price, 3), "50 EMA": ema_50_val or "N/A",
                "200 EMA": ema_200_val or "N/A", "NAV": round(nav, 3), 
                "P/NAV": round(p_nav, 2), "Yield (%)": round(div_yield, 2), 
                "Payout Check": payout_status,
                "Gearing (%)": round(gearing, 2) if gearing_calculated else "N/A",
                "ICR": icr if icr is not None else "N/A",
                "Debt Risk Tier": debt_tier, "Signal": signal
            }
            results.append(reit_metrics)

            # Individual Tab Export
            tab_name = symbol.replace(".SI", "").replace(".HK", "")[:31]
            ws_reit = wb.create_sheet(title=tab_name)
            ws_reit['A1'] = f"{reit_name} ({symbol}) Detailed Analysis"
            ws_reit['A1'].font = Font(size=14, bold=True, color="1F497D")
            ws_reit.append([])
            ws_reit.append(["Metric", "Value"])
            
            for k, v in reit_metrics.items():
                ws_reit.append([k, v])

            if chart_img_path and os.path.exists(chart_img_path):
                img = OpenpyxlImage(chart_img_path)
                ws_reit.add_image(img, 'D3')
                temp_images.append(chart_img_path)

        except Exception as e:
            print(f"Error processing {symbol}: {e}")
            continue

    # --- DASHBOARD GENERATION & FORMATTING ---
    df = pd.DataFrame(results)
    ws_dashboard['A1'] = "REIT HYBRID SCREENER REPORT"
    ws_dashboard['A2'] = f"Generated: {current_time} | Benchmark Rate: {risk_free_rate}% | Target Yield: >{round(dynamic_yield_threshold, 2)}%"
    ws_dashboard['A2'].font = Font(size=10, italic=True, color="595959")
    
    for r in dataframe_to_rows(df, index=False, header=True):
        ws_dashboard.append(r)

    colors = {
        "STRONG ACCUMULATE": ("92D050", "004E00", True),
        "ACCUMULATE": ("C6EFCE", "006100", False),
        "HOLD": ("FFEB9C", "9C6500", False),
        "DECUMULATE": ("FFC7CE", "9C0006", False),
        "STRONG DECUMULATE": ("FF0000", "FFFFFF", True)
    }

    sig_col_idx = chr(ord('A') + list(reit_metrics.keys()).index("Signal"))
    
    for row in range(4, ws_dashboard.max_row + 1):
        cell = ws_dashboard[f"{sig_col_idx}{row}"]
        if cell.value in colors:
            bg, fg, bold = colors[cell.value]
            cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            cell.font = Font(color=fg, bold=bold)

    # --- DISCLAIMER FOOTER ENGINE ---
    disclaimer_row = ws_dashboard.max_row + 3
    disclaimer_text = (
        "Disclaimer: The information expressed is for reference purposes only and does not constitute "
        "any form of investment advice. Please do your own due diligence or speak to a licensed financial "
        "adviser before making any form of investment, as all investment involves risk. "
        "The advertisement has not been reviewed by the Monetary Authority of Singapore."
    )
    
    ws_dashboard.cell(row=disclaimer_row, column=1, value=disclaimer_text)
    disc_cell = ws_dashboard.cell(row=disclaimer_row, column=1)
    disc_cell.font = Font(size=9, italic=True, color="7F7F7F")

    # --- FREEZE PANES ENGINE ---
    ws_dashboard.freeze_panes = 'B4'

    # --- COLUMN SPACING ENGINE ---
    for col in ws_dashboard.columns:
        col_letter = col[0].column_letter
        max_len = 0
        for cell in col[2:ws_dashboard.max_row - 2]: # Ignore titles & disclaimer row length
            cell_len = len(str(cell.value or ''))
            if cell_len > max_len:
                max_len = cell_len
        
        calculated_width = max(max_len + 3, 11)
        ws_dashboard.column_dimensions[col_letter].width = min(calculated_width, 32)

    wb.save(output_filename)
    
    for path in temp_images:
        if os.path.exists(path):
            os.remove(path)
    gc.collect()

if __name__ == "__main__":
    get_reit_data()
